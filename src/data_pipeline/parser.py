"""
Document and Resume Parser
Extracts raw text, contact information, structural sections, and bullet points.
"""

import io
import re
from typing import ClassVar
import xml.etree.ElementTree as ET
import zipfile

import docx
import pypdf


class ResumeParser:
    """Extracts structured fields and sections from PDF or plaintext resumes."""

    SECTION_HEADERS: ClassVar[dict[str, list[str]]] = {
        "summary": [
            r"summary", r"professional summary", r"executive summary",
            r"profile", r"about me", r"objective", r"career objective"
        ],
        "experience": [
            r"experience", r"work experience", r"professional experience",
            r"employment history", r"work history", r"career history"
        ],
        "education": [
            r"education", r"academic background", r"academic qualifications",
            r"qualifications", r"degrees"
        ],
        "skills": [
            r"skills", r"technical skills", r"core competencies",
            r"technologies", r"tools & technologies", r"skills & expertise"
        ],
        "projects": [
            r"projects", r"personal projects", r"key projects",
            r"academic projects", r"open source"
        ],
        "certifications": [
            r"certifications", r"licenses", r"courses", r"credentials",
            r"certificates & awards"
        ]
    }

    EMAIL_PATTERN = re.compile(
        r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+", re.IGNORECASE
    )
    PHONE_PATTERN = re.compile(
        r"(?:(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4})", re.IGNORECASE
    )
    LINKEDIN_PATTERN = re.compile(
        r"(?:https?://)?(?:www\.)?linkedin\.com/in/[a-zA-Z0-9_-]+", re.IGNORECASE
    )
    GITHUB_PATTERN = re.compile(
        r"(?:https?://)?(?:www\.)?github\.com/[a-zA-Z0-9_-]+", re.IGNORECASE
    )

    @classmethod
    def extract_text_from_pdf(cls, file_bytes: bytes) -> str:
        """Extract text content from uploaded PDF bytes."""
        try:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            text_pages = []
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text_pages.append(extracted)
            return "\n\n".join(text_pages)
        except Exception as e:
            return f"Error extracting PDF: {e!s}"

    @classmethod
    def extract_text_from_docx(cls, file_bytes: bytes) -> str:
        """Extract text content from uploaded DOCX bytes."""
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        paragraphs.append(row_text)
                        
            return "\n\n".join(paragraphs)
        except Exception as e:
            # Fallback to direct XML extraction if docx fails
            try:
                with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
                    xml_content = z.read("word/document.xml")
                    tree = ET.fromstring(xml_content)
                    texts = [node.text for node in tree.iter() if node.text]
                    return " ".join(texts)
            except Exception as xml_e:
                return f"Error extracting DOCX: {e!s} / {xml_e!s}"

    @classmethod
    def parse_file(cls, filename: str, file_bytes: bytes) -> str:
        """Unified parser routing based on file extension."""
        name_lower = filename.lower()
        if name_lower.endswith(".pdf"):
            return cls.extract_text_from_pdf(file_bytes)
        elif name_lower.endswith(".docx"):
            return cls.extract_text_from_docx(file_bytes)
        else:
            return file_bytes.decode("utf-8", errors="ignore")

    @classmethod
    def extract_contact_info(cls, text: str) -> dict[str, str | None]:
        """Extract name heuristic, email, phone, and profile links."""
        lines = [line.strip() for line in text.split("\n") if line.strip()]
        
        # Candidate name heuristic: first non-empty line that looks like a name
        candidate_name = None
        for line in lines[:5]:
            clean = re.sub(r"[^a-zA-Z\s]", "", line).strip()
            words = clean.split()
            if (
                1 < len(words) <= 4
                and all(w[0].isupper() for w in words if w)
                and not any(header in clean.lower() for headers in cls.SECTION_HEADERS.values() for header in headers)
            ):
                candidate_name = clean
                break
        
        email_match = cls.EMAIL_PATTERN.search(text)
        phone_match = cls.PHONE_PATTERN.search(text)
        linkedin_match = cls.LINKEDIN_PATTERN.search(text)
        github_match = cls.GITHUB_PATTERN.search(text)

        return {
            "name": candidate_name or "Candidate",
            "email": email_match.group(0) if email_match else None,
            "phone": phone_match.group(0) if phone_match else None,
            "linkedin": linkedin_match.group(0) if linkedin_match else None,
            "github": github_match.group(0) if github_match else None,
        }

    @classmethod
    def segment_sections(cls, text: str) -> dict[str, str]:
        """Segment resume into major sections based on detected headings."""
        lines = text.split("\n")
        section_indices: list[tuple[int, str]] = []

        # Find line positions for each heading
        for i, raw_line in enumerate(lines):
            line = raw_line.strip()
            # Clean punctuation from potential headers
            clean_line = re.sub(r"[:\-_#*]", "", line).strip().lower()
            
            for section_name, patterns in cls.SECTION_HEADERS.items():
                if any(clean_line == p.lower() or re.match(rf"^{p}$", clean_line) for p in patterns):
                    section_indices.append((i, section_name))
                    break

        # Sort indices
        section_indices = sorted(section_indices, key=lambda x: x[0])
        
        sections: dict[str, str] = {
            "header": "",
            "summary": "",
            "experience": "",
            "education": "",
            "skills": "",
            "projects": "",
            "certifications": "",
            "other": ""
        }

        if not section_indices:
            # If no explicit headings, treat whole text as experience/content
            sections["experience"] = text
            return sections

        # Header is everything before the first section heading
        first_idx = section_indices[0][0]
        sections["header"] = "\n".join(lines[:first_idx]).strip()

        for k in range(len(section_indices)):
            start_line, sec_name = section_indices[k]
            end_line = section_indices[k + 1][0] if (k + 1 < len(section_indices)) else len(lines)
            content = "\n".join(lines[start_line + 1:end_line]).strip()
            sections[sec_name] = content

        return sections

    @classmethod
    def extract_bullet_points(cls, text: str) -> list[str]:
        """Extract individual bullet points or meaningful sentences from experience / project text."""
        raw_lines = text.split("\n")
        bullets: list[str] = []
        current_bullet: list[str] = []

        bullet_prefixes = (r"^[•\-\*\u2022\u2023\u25E6\u2043\u2219]\s*", r"^\d+[\.\)]\s*")

        for line in raw_lines:
            stripped = line.strip()
            if not stripped:
                if current_bullet:
                    bullets.append(" ".join(current_bullet).strip())
                    current_bullet = []
                continue

            is_bullet_start = any(re.match(pattern, stripped) for pattern in bullet_prefixes)
            if is_bullet_start:
                if current_bullet:
                    bullets.append(" ".join(current_bullet).strip())
                # Remove the bullet mark
                clean = re.sub(r"^[•\-\*\u2022\u2023\u25E6\u2043\u2219\d\.\)]\s*", "", stripped)
                current_bullet = [clean]
            elif current_bullet:
                # Continuation of current bullet
                current_bullet.append(stripped)
            else:
                # Normal sentence or line
                if len(stripped.split()) >= 5:
                    current_bullet = [stripped]

        if current_bullet:
            bullets.append(" ".join(current_bullet).strip())

        # Clean bullets
        return [b for b in bullets if len(b.split()) >= 4]
