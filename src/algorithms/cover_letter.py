"""
Algorithmic Tailored Cover Letter Generator and DOCX Exporter
Constructs role-tailored cover letters linking candidate achievements to company requirements.
"""

import io
import re
from typing import Dict, List, Optional
import docx
from docx.shared import Inches, Pt, RGBColor


class CoverLetterGenerator:
    """Generates customized, high-converting cover letters based on candidate metrics and job posting."""

    @classmethod
    def generate_cover_letter(
        cls,
        candidate_name: str,
        contact_info: Dict[str, any],
        target_role: str,
        job_text: str,
        top_accomplishments: List[str],
        matched_skills: List[str]
    ) -> str:
        """
        Synthesize a structured 3-part cover letter:
        1. Compelling Hook & Role Alignment
        2. Evidence-Based Achievements & Metric Highlights
        3. Strategic Alignment & Call to Action
        """
        # Heuristic company name extraction from job text
        company_match = re.search(r"(?:at|with|join|about)\s+([A-Z][a-zA-Z0-9&]+(?:\s+[A-Z][a-zA-Z0-9&]+)?(?:\s+(?:Inc|LLC|Technologies|Labs|Systems|AI|Group|Solutions))?)", job_text)
        company_name = company_match.group(1) if company_match else "your organization"

        email = contact_info.get("email") or "email@example.com"
        phone = contact_info.get("phone") or "(555) 000-0000"
        linkedin = contact_info.get("linkedin") or ""

        skills_sample = ", ".join(matched_skills[:5]) if matched_skills else "distributed systems and modern software architecture"
        bullet_1 = top_accomplishments[0] if len(top_accomplishments) > 0 else "Engineered scalable production systems delivering high availability and reduced operational latency."
        bullet_2 = top_accomplishments[1] if len(top_accomplishments) > 1 else "Optimized core algorithmic workflows, saving compute overhead and improving reliability."

        letter = f"""{candidate_name}
{email} | {phone} | {linkedin}

Hiring Team
{company_name}

RE: Application for {target_role}

Dear Hiring Team,

I am writing to express my strong interest in the {target_role} position at {company_name}. With proven experience architecting high-impact solutions across {skills_sample}, I have built my career delivering measurable technical outcomes and driving engineering excellence. Having followed {company_name}'s recent work, I am eager to apply my technical and domain expertise to accelerate your team's roadmap.

Throughout my career, I have prioritized high-throughput engineering coupled with tangible business impact. In my recent roles, I have:
• {bullet_1}
• {bullet_2}

What excites me most about {company_name} is the opportunity to solve complex, high-scale engineering challenges. My background in {skills_sample} positions me to contribute immediately to your production pipelines, collaborate cross-functionally, and uphold rigorous engineering standards.

Thank you for your time and consideration. I welcome the opportunity to discuss how my technical skills and track record of delivery align with {company_name}'s mission.

Sincerely,

{candidate_name}
"""
        return letter.strip()

    @classmethod
    def generate_docx(cls, cover_letter_text: str, candidate_name: str = "Candidate") -> bytes:
        """Export cover letter as a clean, professionally formatted DOCX document."""
        doc = docx.Document()
        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.85)
            section.right_margin = Inches(0.85)

        paragraphs = cover_letter_text.split("\n\n")

        for idx, para_text in enumerate(paragraphs):
            p = doc.add_paragraph()
            lines = para_text.split("\n")

            if idx == 0:  # Header
                for line in lines:
                    r = p.add_run(line + "\n")
                    r.font.name = "Calibri"
                    r.font.size = Pt(10.5)
                    r.font.color.rgb = RGBColor(71, 85, 105)
                p.paragraph_format.space_after = Pt(12)
            else:
                for line in lines:
                    if line.startswith("•"):
                        bp = doc.add_paragraph(style='List Bullet')
                        br = bp.add_run(line.replace("•", "").strip())
                        br.font.name = "Calibri"
                        br.font.size = Pt(11)
                        bp.paragraph_format.space_after = Pt(3)
                    else:
                        r = p.add_run(line + "\n")
                        r.font.name = "Calibri"
                        r.font.size = Pt(11)
                        r.font.color.rgb = RGBColor(30, 41, 59)
                        p.paragraph_format.space_after = Pt(8)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
