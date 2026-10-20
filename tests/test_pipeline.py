"""
Unit Tests for the AdaptiveML Resume and Machine Learning Pipeline
"""

import unittest
from src.algorithms.ats_scorer import ATSScorer
from src.algorithms.matcher import JobFitMatcher
from src.algorithms.qa_engine import AdaptiveQAEngine
from src.algorithms.text_rank import TextRankSummarizer
from src.data_pipeline.parser import ResumeParser
from src.data_pipeline.preprocessor import ResumePreprocessor
from src.data_pipeline.skill_taxonomy import extract_skills_from_text
from src.feature_engineering.embeddings import DenseEmbedder
from src.feature_engineering.vectorizers import BM25Okapi, TFIDFEngine
from src.models.cluster_engine import ExperienceClusterEngine
from src.models.domain_classifier import DomainClassifier
from src.models.intent_classifier import AdaptiveIntentClassifier
from src.models.topic_modeler import TopicModeler
from src.sample_data import SAMPLE_JOB_DESCRIPTIONS, SAMPLE_RESUMES


class TestAdaptiveMLPipeline(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.resume_text = SAMPLE_RESUMES["Senior Machine Learning Engineer (Alex Rivera)"]
        cls.job_text = SAMPLE_JOB_DESCRIPTIONS["Staff / Senior Machine Learning Engineer"]
        cls.embedder = DenseEmbedder(use_transformer=False) # Test high-speed LSA mode
        cls.domain_clf = DomainClassifier()
        cls.intent_clf = AdaptiveIntentClassifier()
        cls.matcher = JobFitMatcher(cls.embedder)
        cls.summarizer = TextRankSummarizer()
        cls.qa_engine = AdaptiveQAEngine(
            intent_classifier=cls.intent_clf,
            domain_classifier=cls.domain_clf,
            embedder=cls.embedder,
            matcher=cls.matcher,
            summarizer=cls.summarizer
        )

    def test_parser_contact_info(self):
        contact = ResumeParser.extract_contact_info(self.resume_text)
        self.assertEqual(contact["name"], "Alex Rivera")
        self.assertEqual(contact["email"], "alex.rivera@email.com")
        self.assertIn("555", contact["phone"])
        self.assertIn("linkedin.com", contact["linkedin"])

    def test_parser_sections_and_bullets(self):
        sections = ResumeParser.segment_sections(self.resume_text)
        self.assertTrue(len(sections["experience"]) > 0)
        self.assertTrue(len(sections["skills"]) > 0)
        
        bullets = ResumeParser.extract_bullet_points(self.resume_text)
        self.assertGreaterEqual(len(bullets), 4)

    def test_skill_taxonomy_extraction(self):
        skills = extract_skills_from_text(self.resume_text)
        flat_skills = [s.lower() for cat in skills.values() for s in cat]
        self.assertIn("python", flat_skills)
        self.assertIn("pytorch", flat_skills)
        self.assertIn("docker", flat_skills)

    def test_preprocessor_readability_and_bullets(self):
        readability = ResumePreprocessor.calculate_readability(self.resume_text)
        self.assertGreater(readability["flesch_reading_ease"], 0)
        self.assertGreater(readability["total_words"], 100)

        bullets = ResumeParser.extract_bullet_points(self.resume_text)
        bullet_analysis = ResumePreprocessor.analyze_bullets(bullets)
        self.assertGreater(bullet_analysis["action_verb_count"], 0)
        self.assertGreater(bullet_analysis["quantifiable_count"], 0)

    def test_bm25_okapi(self):
        corpus = [
            "Architected and deployed an enterprise Retrieval-Augmented Generation RAG platform",
            "Refactored legacy PostgreSQL queries and implemented Redis caching layers",
            "Orchestrated automated end-to-end testing with Pytest"
        ]
        bm25 = BM25Okapi(corpus)
        scores = bm25.get_scores("RAG Retrieval enterprise")
        self.assertGreater(scores[0], scores[1])

    def test_domain_classifier(self):
        result = self.domain_clf.predict(self.resume_text)
        self.assertEqual(result["predicted_domain"], "Machine Learning & AI")
        self.assertGreater(result["confidence"], 25.0)

    def test_adaptive_intent_classifier(self):
        q1 = self.intent_clf.classify_intent("What is my ATS score?")
        self.assertEqual(q1["intent"], "ATS_DIAGNOSTICS")

        q2 = self.intent_clf.classify_intent("Generate technical interview questions")
        self.assertEqual(q2["intent"], "INTERVIEW_PREP")

    def test_ats_scorer(self):
        contact = ResumeParser.extract_contact_info(self.resume_text)
        sections = ResumeParser.segment_sections(self.resume_text)
        bullets = ResumeParser.extract_bullet_points(self.resume_text)
        skills = extract_skills_from_text(self.resume_text)
        readability = ResumePreprocessor.calculate_readability(self.resume_text)
        bullet_analysis = ResumePreprocessor.analyze_bullets(bullets)

        result = ATSScorer.evaluate(contact, sections, bullet_analysis, skills, readability)
        self.assertGreaterEqual(result["total_score"], 75.0)
        self.assertIn("A", result["grade"])

    def test_job_fit_matcher(self):
        fit = self.matcher.evaluate_fit(self.resume_text, self.job_text)
        self.assertGreater(fit["composite_score"], 50.0)
        self.assertTrue(len(fit["matched_skills"]) > 0)

    def test_text_rank_summarizer(self):
        bullets = ResumeParser.extract_bullet_points(self.resume_text)
        summary = self.summarizer.summarize(bullets, top_n=2)
        self.assertEqual(len(summary), 2)

    def test_qa_engine(self):
        response = self.qa_engine.answer_query(
            query="What is my ATS score?",
            resume_text=self.resume_text
        )
        self.assertEqual(response["intent"], "ATS_DIAGNOSTICS")
        self.assertIn("ATS Score", response["answer"])

    def test_docx_parser(self):
        import docx
        import io
        doc = docx.Document()
        doc.add_heading("Alex Rivera", 0)
        doc.add_paragraph("Senior ML Engineer with PyTorch and Python experience.")
        bio = io.BytesIO()
        doc.save(bio)
        text = ResumeParser.extract_text_from_docx(bio.getvalue())
        self.assertIn("Alex Rivera", text)
        self.assertIn("PyTorch", text)

    def test_resume_tailor(self):
        from src.algorithms.tailor import ResumeTailor
        fit = self.matcher.evaluate_fit(self.resume_text, self.job_text)
        tailored = ResumeTailor.tailor_resume(self.resume_text, fit)
        self.assertGreaterEqual(tailored["ats_after"], tailored["ats_before"])
        docx_bytes = ResumeTailor.generate_docx(tailored["tailored_text"], "Alex Rivera")
        self.assertGreater(len(docx_bytes), 1000)

    def test_cover_letter_generator(self):
        from src.algorithms.cover_letter import CoverLetterGenerator
        bullets = ResumeParser.extract_bullet_points(self.resume_text)
        top_acc = [b[2] for b in self.summarizer.summarize(bullets, top_n=2)]
        contact = ResumeParser.extract_contact_info(self.resume_text)
        letter = CoverLetterGenerator.generate_cover_letter(
            candidate_name="Alex Rivera",
            contact_info=contact,
            target_role="Staff ML Engineer",
            job_text=self.job_text,
            top_accomplishments=top_acc,
            matched_skills=["PyTorch", "Transformers"]
        )
        self.assertIn("Alex Rivera", letter)
        self.assertIn("Staff ML Engineer", letter)
        docx_bytes = CoverLetterGenerator.generate_docx(letter, "Alex Rivera")
        self.assertGreater(len(docx_bytes), 1000)

    def test_seniority_estimator(self):
        from src.algorithms.seniority_estimator import SeniorityEstimator
        bullets = ResumeParser.extract_bullet_points(self.resume_text)
        ba = ResumePreprocessor.analyze_bullets(bullets)
        seniority = SeniorityEstimator.estimate_seniority(self.resume_text, "Machine Learning & AI", ba)
        self.assertIn(seniority["seniority_level"], ["Senior", "Staff / Lead", "Principal"])
        self.assertIn("$", seniority["salary_benchmark"])
        self.assertTrue(len(seniority["technical_milestones"]) > 0)


if __name__ == "__main__":
    unittest.main()
